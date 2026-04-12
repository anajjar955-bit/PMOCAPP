from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from bson import ObjectId
import os
import logging
import bcrypt
import jwt
import secrets
import uuid
from datetime import datetime, timezone, timedelta
from pydantic import BaseModel
from typing import Optional
import pathlib
import base64
from fastapi.responses import Response

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ.get('DB_NAME', 'pmhouse_academy')]

JWT_SECRET = os.environ.get('JWT_SECRET', secrets.token_hex(32))
JWT_ALGORITHM = "HS256"
PAYPAL_LINK = "https://www.paypal.com/ncp/payment/QEL5ME5XAAD96"

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ========== Pydantic Models ==========
class RegisterRequest(BaseModel):
    email: str
    password: str
    name: str

class LoginRequest(BaseModel):
    email: str
    password: str

class QuizSubmission(BaseModel):
    answers: dict

class ExamSubmission(BaseModel):
    answers: dict

class PaymentConfirmation(BaseModel):
    transaction_id: Optional[str] = None

# ========== Auth Helpers ==========
def hash_password(password: str) -> str:
    salt = bcrypt.gensalt()
    return bcrypt.hashpw(password.encode("utf-8"), salt).decode("utf-8")

def verify_password(plain_password: str, hashed_password: str) -> bool:
    return bcrypt.checkpw(plain_password.encode("utf-8"), hashed_password.encode("utf-8"))

def create_access_token(user_id: str, email: str) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "exp": datetime.now(timezone.utc) + timedelta(days=7),
        "type": "access"
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

async def get_current_user(request: Request) -> dict:
    auth_header = request.headers.get("Authorization", "")
    if not auth_header.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Not authenticated")
    token = auth_header[7:]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Invalid token type")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        user["id"] = str(user["_id"])
        del user["_id"]
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")

async def get_optional_user(request: Request):
    try:
        return await get_current_user(request)
    except Exception:
        return None

# ========== Auth Routes ==========
@api_router.post("/auth/register")
async def register(req: RegisterRequest):
    email = req.email.lower().strip()
    if not email or not req.password or len(req.password) < 6:
        raise HTTPException(status_code=400, detail="Invalid email or password (min 6 chars)")
    existing = await db.users.find_one({"email": email})
    if existing:
        raise HTTPException(status_code=400, detail="البريد الإلكتروني مسجل مسبقاً")
    user_doc = {
        "email": email,
        "password_hash": hash_password(req.password),
        "name": req.name,
        "role": "user",
        "is_paid": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "progress": {"completed_lessons": [], "quiz_scores": {}, "exam_attempts": []}
    }
    result = await db.users.insert_one(user_doc)
    user_id = str(result.inserted_id)
    token = create_access_token(user_id, email)
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": user_id, "email": email, "name": req.name,
            "role": "user", "is_paid": False,
            "progress": user_doc["progress"]
        }
    }

@api_router.post("/auth/login")
async def login(req: LoginRequest):
    email = req.email.lower().strip()
    user = await db.users.find_one({"email": email})
    if not user:
        raise HTTPException(status_code=401, detail="بريد إلكتروني أو كلمة مرور غير صحيحة")
    if not verify_password(req.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="بريد إلكتروني أو كلمة مرور غير صحيحة")
    user_id = str(user["_id"])
    token = create_access_token(user_id, email)
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": user_id, "email": user["email"], "name": user.get("name", ""),
            "role": user.get("role", "user"), "is_paid": user.get("is_paid", False),
            "progress": user.get("progress", {"completed_lessons": [], "quiz_scores": {}, "exam_attempts": []})
        }
    }

@api_router.get("/auth/me")
async def get_me(request: Request):
    user = await get_current_user(request)
    return {"user": user}

# ========== Password Reset ==========
class ResetRequest(BaseModel):
    email: str

class ResetConfirm(BaseModel):
    email: str
    code: str
    new_password: str

@api_router.post("/auth/request-reset")
async def request_password_reset(req: ResetRequest):
    email = req.email.lower().strip()
    user = await db.users.find_one({"email": email})
    if not user:
        raise HTTPException(status_code=404, detail="البريد الإلكتروني غير مسجل")
    code = ''.join(random.choices(string.digits, k=6))
    await db.password_resets.update_one(
        {"email": email},
        {"$set": {
            "email": email,
            "code": code,
            "name": user.get("name", ""),
            "used": False,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "expires_at": (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat()
        }},
        upsert=True
    )
    return {"success": True, "message": "تم إنشاء كود إعادة التعيين. تواصل مع الأدمن عبر واتساب لاستلام الكود."}

@api_router.post("/auth/reset-password")
async def reset_password(req: ResetConfirm):
    email = req.email.lower().strip()
    if len(req.new_password) < 6:
        raise HTTPException(status_code=400, detail="كلمة المرور يجب أن تكون 6 أحرف على الأقل")
    record = await db.password_resets.find_one({"email": email, "code": req.code, "used": False}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=400, detail="كود إعادة التعيين غير صحيح أو منتهي الصلاحية")
    if record.get("expires_at") and record["expires_at"] < datetime.now(timezone.utc).isoformat():
        raise HTTPException(status_code=400, detail="انتهت صلاحية الكود. اطلب كود جديد.")
    await db.password_resets.update_one({"email": email, "code": req.code}, {"$set": {"used": True}})
    await db.users.update_one({"email": email}, {"$set": {"password_hash": hash_password(req.new_password)}})
    return {"success": True, "message": "تم تغيير كلمة المرور بنجاح!"}

@api_router.get("/admin/reset-requests")
async def list_reset_requests(request: Request):
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    resets = await db.password_resets.find({}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return {"resets": resets}

# ========== Course Routes ==========
@api_router.get("/course/modules")
async def get_modules(request: Request):
    user = await get_optional_user(request)
    modules = await db.modules.find({}, {"_id": 0}).sort("order", 1).to_list(100)
    for module in modules:
        lessons = await db.lessons.find({"module_id": module["id"]}, {"_id": 0, "id": 1, "is_demo": 1}).to_list(100)
        module["lesson_count"] = len(lessons)
        module["demo_lessons"] = sum(1 for l in lessons if l.get("is_demo", False))
        if user:
            completed = user.get("progress", {}).get("completed_lessons", [])
            module["completed_lessons"] = sum(1 for l in lessons if l["id"] in completed)
        else:
            module["completed_lessons"] = 0
    return {"modules": modules}

@api_router.get("/course/modules/{module_id}/lessons")
async def get_module_lessons(module_id: str, request: Request):
    user = await get_optional_user(request)
    lessons = await db.lessons.find({"module_id": module_id}, {"_id": 0, "slides": 0}).sort("order", 1).to_list(100)
    if user:
        completed = user.get("progress", {}).get("completed_lessons", [])
        quiz_scores = user.get("progress", {}).get("quiz_scores", {})
        is_paid = user.get("is_paid", False)
        for lesson in lessons:
            lesson["is_completed"] = lesson["id"] in completed
            lesson["quiz_score"] = quiz_scores.get(lesson["id"])
            lesson["is_accessible"] = lesson.get("is_demo", False) or is_paid
    else:
        for lesson in lessons:
            lesson["is_completed"] = False
            lesson["quiz_score"] = None
            lesson["is_accessible"] = lesson.get("is_demo", False)
    return {"lessons": lessons}

@api_router.get("/course/lessons/{lesson_id}")
async def get_lesson(lesson_id: str, request: Request):
    user = await get_optional_user(request)
    is_paid = user.get("is_paid", False) if user else False
    lesson = await db.lessons.find_one({"id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="الدرس غير موجود")
    if not lesson.get("is_demo", False) and not is_paid:
        raise HTTPException(status_code=403, detail="يجب الاشتراك للوصول لهذا الدرس")
    return {"lesson": lesson}

@api_router.get("/course/lessons/{lesson_id}/quiz")
async def get_lesson_quiz(lesson_id: str, request: Request):
    user = await get_current_user(request)
    questions = await db.questions.find(
        {"lesson_id": lesson_id}, {"_id": 0, "correct_answer": 0, "explanation": 0}
    ).to_list(100)
    return {"questions": questions}

@api_router.post("/course/lessons/{lesson_id}/quiz/submit")
async def submit_quiz(lesson_id: str, submission: QuizSubmission, request: Request):
    user = await get_current_user(request)
    questions = await db.questions.find({"lesson_id": lesson_id}, {"_id": 0}).to_list(100)
    total = len(questions)
    correct = 0
    results = {}
    for q in questions:
        user_answer = submission.answers.get(q["id"])
        is_correct = user_answer == q["correct_answer"]
        if is_correct:
            correct += 1
        results[q["id"]] = {
            "correct": is_correct,
            "correct_answer": q["correct_answer"],
            "user_answer": user_answer,
            "explanation": q.get("explanation", "")
        }
    score = round((correct / total) * 100) if total > 0 else 0
    await db.users.update_one(
        {"_id": ObjectId(user["id"])},
        {"$set": {f"progress.quiz_scores.{lesson_id}": score}}
    )
    return {"score": score, "correct": correct, "total": total, "results": results}

@api_router.post("/course/lessons/{lesson_id}/complete")
async def complete_lesson(lesson_id: str, request: Request):
    user = await get_current_user(request)
    await db.users.update_one(
        {"_id": ObjectId(user["id"])},
        {"$addToSet": {"progress.completed_lessons": lesson_id}}
    )
    return {"success": True}

@api_router.get("/course/lessons/{lesson_id}/next")
async def get_next_lesson(lesson_id: str, request: Request):
    """Get the next lesson after the current one"""
    user = await get_optional_user(request)
    is_paid = user.get("is_paid", False) if user else False
    current = await db.lessons.find_one({"id": lesson_id}, {"_id": 0, "module_id": 1, "order": 1})
    if not current:
        return {"next_lesson": None}
    next_in_module = await db.lessons.find_one(
        {"module_id": current["module_id"], "order": current["order"] + 1},
        {"_id": 0, "id": 1, "title": 1, "is_demo": 1, "module_id": 1}
    )
    if next_in_module:
        next_in_module["is_accessible"] = next_in_module.get("is_demo", False) or is_paid
        return {"next_lesson": next_in_module}
    current_mod = await db.modules.find_one({"id": current["module_id"]}, {"_id": 0, "order": 1})
    if current_mod:
        next_mod = await db.modules.find_one({"order": current_mod["order"] + 1}, {"_id": 0, "id": 1})
        if next_mod:
            first_lesson = await db.lessons.find_one(
                {"module_id": next_mod["id"], "order": 1},
                {"_id": 0, "id": 1, "title": 1, "is_demo": 1, "module_id": 1}
            )
            if first_lesson:
                first_lesson["is_accessible"] = first_lesson.get("is_demo", False) or is_paid
                return {"next_lesson": first_lesson}
    return {"next_lesson": None}

# ========== Exam Routes ==========
@api_router.get("/exams")
async def get_exams(request: Request):
    user = await get_optional_user(request)
    exams = await db.exams.find({}, {"_id": 0}).to_list(10)
    if user:
        for exam in exams:
            attempts = await db.exam_attempts.find(
                {"user_id": user["id"], "exam_id": exam["id"]},
                {"_id": 0, "score": 1, "created_at": 1, "correct": 1, "total": 1}
            ).sort("created_at", -1).to_list(10)
            exam["attempts"] = attempts
            exam["is_accessible"] = user.get("is_paid", False)
    else:
        for exam in exams:
            exam["attempts"] = []
            exam["is_accessible"] = False
    return {"exams": exams}

@api_router.get("/exams/{exam_id}")
async def get_exam(exam_id: str, request: Request):
    user = await get_current_user(request)
    if not user.get("is_paid", False):
        raise HTTPException(status_code=403, detail="يجب الاشتراك للوصول للاختبارات")
    exam = await db.exams.find_one({"id": exam_id}, {"_id": 0})
    if not exam:
        raise HTTPException(status_code=404, detail="الاختبار غير موجود")
    questions = await db.exam_questions.find(
        {"exam_id": exam_id}, {"_id": 0, "correct_answer": 0, "explanation": 0}
    ).to_list(100)
    return {"exam": exam, "questions": questions}

@api_router.post("/exams/{exam_id}/submit")
async def submit_exam(exam_id: str, submission: ExamSubmission, request: Request):
    user = await get_current_user(request)
    questions = await db.exam_questions.find({"exam_id": exam_id}, {"_id": 0}).to_list(100)
    total = len(questions)
    correct = 0
    results = {}
    domain_scores = {}
    for q in questions:
        user_answer = submission.answers.get(q["id"])
        is_correct = user_answer == q["correct_answer"]
        if is_correct:
            correct += 1
        domain = q.get("domain", "unknown")
        if domain not in domain_scores:
            domain_scores[domain] = {"correct": 0, "total": 0}
        domain_scores[domain]["total"] += 1
        if is_correct:
            domain_scores[domain]["correct"] += 1
        results[q["id"]] = {
            "correct": is_correct,
            "correct_answer": q["correct_answer"],
            "user_answer": user_answer,
            "explanation": q.get("explanation", "")
        }
    score = round((correct / total) * 100) if total > 0 else 0
    attempt = {
        "id": str(uuid.uuid4()),
        "user_id": user["id"],
        "exam_id": exam_id,
        "score": score,
        "correct": correct,
        "total": total,
        "domain_scores": domain_scores,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.exam_attempts.insert_one(attempt)
    await db.users.update_one(
        {"_id": ObjectId(user["id"])},
        {"$push": {"progress.exam_attempts": {
            "exam_id": exam_id, "score": score,
            "date": datetime.now(timezone.utc).isoformat()
        }}}
    )
    return {"score": score, "correct": correct, "total": total, "results": results, "domain_scores": domain_scores}

# ========== Progress Routes ==========
@api_router.get("/progress")
async def get_progress(request: Request):
    user = await get_current_user(request)
    total_lessons = await db.lessons.count_documents({})
    completed = len(user.get("progress", {}).get("completed_lessons", []))
    return {
        "total_lessons": total_lessons,
        "completed_lessons": completed,
        "completed_lesson_ids": user.get("progress", {}).get("completed_lessons", []),
        "progress_percentage": round((completed / total_lessons) * 100) if total_lessons > 0 else 0,
        "quiz_scores": user.get("progress", {}).get("quiz_scores", {}),
        "exam_attempts": user.get("progress", {}).get("exam_attempts", []),
        "is_course_complete": completed >= total_lessons
    }

# ========== Audio TTS Routes (ElevenLabs) ==========
@api_router.get("/audio/slide/{lesson_id}/{slide_index}")
async def get_slide_audio(lesson_id: str, slide_index: int, request: Request):
    """Generate TTS audio with ElevenLabs Egyptian Arabic voice"""
    cache_key = f"{lesson_id}_{slide_index}_el1"
    cached = await db.audio_cache.find_one({"cache_key": cache_key}, {"_id": 0})
    if cached and cached.get("audio_base64"):
        audio_bytes = base64.b64decode(cached["audio_base64"])
        return Response(content=audio_bytes, media_type="audio/mp3",
                       headers={"Cache-Control": "public, max-age=86400"})

    lesson = await db.lessons.find_one({"id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Lesson not found")
    slides = lesson.get("slides", [])
    if slide_index < 0 or slide_index >= len(slides):
        raise HTTPException(status_code=404, detail="Slide not found")

    slide = slides[slide_index]
    slide_content = f"عنوان: {slide['title']}\nمحتوى: {slide.get('content', '')}\n"
    key_points = slide.get('key_points', [])
    if key_points:
        slide_content += "نقاط: " + " / ".join(key_points)

    # Step 1: Generate Egyptian Arabic narration using GPT
    try:
        from emergentintegrations.llm.openai import LlmChat, UserMessage
        import hashlib
        sid = f"el_{hashlib.md5(cache_key.encode()).hexdigest()[:8]}"
        chat = LlmChat(
            api_key=os.getenv("EMERGENT_LLM_KEY"),
            session_id=sid,
            system_message="أنت مدرب مصري متخصص في إدارة المشاريع. بتشرح بالعامية المصرية الطبيعية زي ما بتتكلم مع صحابك."
        )
        chat = chat.with_model("openai", "gpt-4o-mini")
        prompt = f"""اشرح المحتوى ده كأنك مدرب مصري بتشرح في فيديو:

{slide_content}

اكتب شرح بالعامية المصرية الطبيعية في حدود 100-130 كلمة.
- ابدأ بمقدمة حماسية قصيرة
- اشرح بهدوء مع أمثلة من الواقع
- اختم بخلاصة سريعة
- استخدم كلمات مصرية: دلوقتي، كده، عشان، يعني، بتاع، حاجة، إزاي، ليه
- اكتب كلام متصل بدون عناوين أو نقاط
- خلي الكلام طبيعي زي حد بيتكلم مش بيقرأ
- مهم جداً: ضع تشكيل كامل على كل الكلمات العربية (فتحة، ضمة، كسرة، سكون، شدة، تنوين) عشان القراءة تكون مظبوطة 100%"""

        narration = await chat.send_message(UserMessage(text=prompt))
        narration = narration.strip()[:4096]
        logger.info(f"Narration for {cache_key}: {narration[:80]}...")
    except Exception as e:
        logger.error(f"GPT narration failed: {e}")
        narration = f"{slide['title']}. {slide.get('content', '')}."

    # Step 2: Generate audio using ElevenLabs with Egyptian voice
    try:
        from elevenlabs import ElevenLabs as ElevenLabsClient
        from elevenlabs.types import VoiceSettings
        el_client = ElevenLabsClient(api_key=os.getenv("ELEVENLABS_API_KEY"))
        voice_id = os.getenv("ELEVENLABS_VOICE_ID", "rMheqEfwsIJckq2yCdb5")

        audio_gen = el_client.text_to_speech.convert(
            text=narration,
            voice_id=voice_id,
            model_id="eleven_multilingual_v2",
            voice_settings=VoiceSettings(
                stability=0.35,
                similarity_boost=0.85,
                style=0.7,
                use_speaker_boost=True
            )
        )
        audio_data = b""
        for chunk in audio_gen:
            audio_data += chunk

        # Cache
        audio_b64 = base64.b64encode(audio_data).decode("utf-8")
        await db.audio_cache.update_one(
            {"cache_key": cache_key},
            {"$set": {"cache_key": cache_key, "audio_base64": audio_b64,
                      "narration_text": narration,
                      "created_at": datetime.now(timezone.utc).isoformat()}},
            upsert=True
        )
        return Response(content=audio_data, media_type="audio/mp3",
                       headers={"Cache-Control": "public, max-age=86400"})
    except Exception as e:
        logger.error(f"ElevenLabs TTS failed: {e}")
        raise HTTPException(status_code=500, detail=f"فشل توليد الصوت: {str(e)}")

# ========== Activation Code System ==========
import random, string

def generate_activation_code():
    chars = string.ascii_uppercase + string.digits
    part1 = ''.join(random.choices(chars, k=4))
    part2 = ''.join(random.choices(chars, k=4))
    return f"PMH-{part1}-{part2}"

@api_router.post("/admin/generate-code")
async def generate_code(request: Request):
    """Admin generates an activation code"""
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    body = await request.json()
    email = body.get("email", "").lower().strip()
    code = generate_activation_code()
    await db.activation_codes.update_one(
        {"email": email},
        {"$set": {
            "email": email,
            "code": code,
            "used": False,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "created_by": user["id"]
        }},
        upsert=True
    )
    return {"code": code, "email": email}

@api_router.post("/activate")
async def activate_with_code(request: Request):
    """User enters activation code to unlock full access"""
    user = await get_current_user(request)
    body = await request.json()
    code = body.get("code", "").strip().upper()
    if not code:
        raise HTTPException(status_code=400, detail="يرجى إدخال كود التفعيل")
    record = await db.activation_codes.find_one({"code": code, "used": False}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=400, detail="كود التفعيل غير صحيح أو مستخدم مسبقاً")
    await db.activation_codes.update_one(
        {"code": code},
        {"$set": {"used": True, "used_by": user["id"], "used_at": datetime.now(timezone.utc).isoformat()}}
    )
    await db.users.update_one(
        {"_id": ObjectId(user["id"])},
        {"$set": {"is_paid": True, "payment_date": datetime.now(timezone.utc).isoformat(), "activation_code": code}}
    )
    return {"success": True, "message": "تم تفعيل اشتراكك بنجاح! 🎉"}

@api_router.get("/admin/codes")
async def list_codes(request: Request):
    """Admin lists all activation codes"""
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    codes = await db.activation_codes.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)
    return {"codes": codes}

@api_router.get("/admin/users")
async def list_users(request: Request):
    """Admin lists all users with progress"""
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    total_lessons = await db.lessons.count_documents({})
    users_cursor = db.users.find({}, {"_id": 0, "password_hash": 0}).sort("created_at", -1)
    users = await users_cursor.to_list(500)
    result = []
    for u in users:
        completed = len(u.get("progress", {}).get("completed_lessons", []))
        quiz_scores = u.get("progress", {}).get("quiz_scores", {})
        avg_quiz = round(sum(quiz_scores.values()) / len(quiz_scores)) if quiz_scores else 0
        exam_attempts = u.get("progress", {}).get("exam_attempts", [])
        best_exam = max([a.get("score", 0) for a in exam_attempts], default=0) if exam_attempts else 0
        result.append({
            "name": u.get("name", ""),
            "email": u.get("email", ""),
            "role": u.get("role", "user"),
            "is_paid": u.get("is_paid", False),
            "created_at": u.get("created_at", ""),
            "completed_lessons": completed,
            "total_lessons": total_lessons,
            "progress_pct": round((completed / total_lessons) * 100) if total_lessons > 0 else 0,
            "avg_quiz_score": avg_quiz,
            "exam_attempts_count": len(exam_attempts),
            "best_exam_score": best_exam,
            "activation_code": u.get("activation_code", ""),
        })
    return {"users": result, "total": len(result)}

@api_router.get("/admin/export-users")
async def export_users_excel(request: Request):
    """Admin exports users data to Excel"""
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    total_lessons = await db.lessons.count_documents({})
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).sort("created_at", -1).to_list(500)
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.sheet_view.rightToLeft = True
    headers = ["الاسم", "البريد الإلكتروني", "الحالة", "مشترك", "تاريخ التسجيل",
               "الدروس المكتملة", "نسبة التقدم", "متوسط الاختبارات القصيرة",
               "عدد محاولات الامتحان", "أعلى درجة امتحان", "كود التفعيل"]
    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, u in enumerate(users, 2):
        completed = len(u.get("progress", {}).get("completed_lessons", []))
        quiz_scores = u.get("progress", {}).get("quiz_scores", {})
        avg_quiz = round(sum(quiz_scores.values()) / len(quiz_scores)) if quiz_scores else 0
        exam_attempts = u.get("progress", {}).get("exam_attempts", [])
        best_exam = max([a.get("score", 0) for a in exam_attempts], default=0) if exam_attempts else 0
        pct = round((completed / total_lessons) * 100) if total_lessons > 0 else 0
        ws.cell(row=row_idx, column=1, value=u.get("name", ""))
        ws.cell(row=row_idx, column=2, value=u.get("email", ""))
        ws.cell(row=row_idx, column=3, value=u.get("role", "user"))
        ws.cell(row=row_idx, column=4, value="نعم" if u.get("is_paid") else "لا")
        ws.cell(row=row_idx, column=5, value=str(u.get("created_at", ""))[:10])
        ws.cell(row=row_idx, column=6, value=f"{completed}/{total_lessons}")
        ws.cell(row=row_idx, column=7, value=f"{pct}%")
        ws.cell(row=row_idx, column=8, value=f"{avg_quiz}%")
        ws.cell(row=row_idx, column=9, value=len(exam_attempts))
        ws.cell(row=row_idx, column=10, value=f"{best_exam}%")
        ws.cell(row=row_idx, column=11, value=u.get("activation_code", ""))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=pmhouse_users.xlsx"}
    )

# ========== Payment Routes ==========
WHATSAPP_NUMBER = "201005394312"

@api_router.get("/payment/info")
async def get_payment_info():
    return {
        "paypal_link": PAYPAL_LINK,
        "whatsapp_number": WHATSAPP_NUMBER,
        "whatsapp_link": f"https://wa.me/{WHATSAPP_NUMBER}"
    }

# ========== Certificate Routes ==========
@api_router.get("/certificate")
async def get_certificate(request: Request):
    user = await get_current_user(request)
    total_lessons = await db.lessons.count_documents({})
    completed = len(user.get("progress", {}).get("completed_lessons", []))
    if completed < total_lessons:
        raise HTTPException(status_code=400, detail=f"الدورة غير مكتملة. أكملت {completed} من {total_lessons} درس")
    return {
        "certificate": {
            "name": user.get("name", ""),
            "email": user.get("email", ""),
            "course_name": "PMI-PMO CP Exam Preparation",
            "course_name_ar": "الإعداد لاختبار شهادة محترف مكتب إدارة المشاريع",
            "issuer": "PM House",
            "issue_date": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
            "certificate_id": f"PMHOUSE-{user['id'][:8].upper()}"
        }
    }

# ========== Admin Routes ==========
@api_router.post("/admin/activate-user")
async def activate_user(request: Request):
    user = await get_current_user(request)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    body = await request.json()
    email = body.get("email", "").lower().strip()
    result = await db.users.update_one(
        {"email": email},
        {"$set": {"is_paid": True, "payment_date": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    return {"success": True}

# ========== Include Router & Middleware ==========
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ========== Startup ==========
@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    # Seed admin
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@pmhouse.com")
    admin_password = os.environ.get("ADMIN_PASSWORD", "PMHouse@2024")
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "email": admin_email,
            "password_hash": hash_password(admin_password),
            "name": "Admin",
            "role": "admin",
            "is_paid": True,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "progress": {"completed_lessons": [], "quiz_scores": {}, "exam_attempts": []}
        })
        logger.info(f"Admin user created: {admin_email}")
    # Seed course data
    from course_data import get_course_data
    module_count = await db.modules.count_documents({})
    if module_count == 0:
        modules, lessons, questions, exams, exam_questions = get_course_data()
        if modules:
            await db.modules.insert_many(modules)
        if lessons:
            await db.lessons.insert_many(lessons)
        if questions:
            await db.questions.insert_many(questions)
        if exams:
            await db.exams.insert_many(exams)
        if exam_questions:
            await db.exam_questions.insert_many(exam_questions)
        logger.info("Course data seeded successfully!")
    # Write test credentials
    memory_dir = pathlib.Path("/app/memory")
    memory_dir.mkdir(exist_ok=True)
    with open("/app/memory/test_credentials.md", "w") as f:
        f.write(f"# Test Credentials\n\n")
        f.write(f"## Admin\n- Email: {admin_email}\n- Password: {admin_password}\n- Role: admin\n\n")
        f.write(f"## Test User\n- Register with any email/password (min 6 chars)\n\n")
        f.write(f"## Auth Endpoints\n- POST /api/auth/register\n- POST /api/auth/login\n- GET /api/auth/me\n\n")
        f.write(f"## PayPal Link\n- {PAYPAL_LINK}\n")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
